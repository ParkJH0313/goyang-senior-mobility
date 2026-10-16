# -*- coding: utf-8 -*-
"""병원·진료과목 원자료와 SGIS 경계로 의료기관 CSV 하나를 생성한다.

선정: 고양시+주변 12개 지역, 내과 또는 정형외과 전문의 >= 1,
      공백을 제거한 기관명에 '치과' 또는 '요양병원'이 없는 기관.
외래 제공·운영시간·휴폐업 여부는 별도 판정하지 않는다.
누락 좌표는 Kakao 주소 API로 보완하고 응답 원문을 raw에 보관한다.
API 키: 프로젝트 루트 .env 또는 환경변수 KAKAO_API_KEY. 코드에 키를 넣지 않는다.
실행 결과·검증은 JSON 표준출력으로만 반환하며 MD·이미지는 생성하지 않는다.
"""
from __future__ import annotations

import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.preprocessing_common import (ROOT, initialize, internal, data_path, save_report,
    building_constants, validate_buildings, building_application_log, walk_constants, validate_walk)
initialize()
from src.preprocessing_common import validate_geocodes

import argparse
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import geopandas as gpd
import pandas as pd
import requests
import shapely
from dotenv import load_dotenv
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from tqdm import tqdm
from urllib3.util.retry import Retry

ID = "암호화요양기호"
ENDPOINT = "https://dapi.kakao.com/v2/local/search/address.json"
DEPARTMENTS = {"01": "내과", "05": "정형외과"}
# HIRA 코드와 SGIS 코드는 서로 다른 체계다. 이름까지 대조한다.
# HIRA: (시도명, 시군구명, SGIS 코드, SGIS 명칭)
REGIONS = {
    "311901": ("경기", "고양덕양구", "31101", "고양시 덕양구"),
    "311903": ("경기", "고양일산동구", "31103", "고양시 일산동구"),
    "311902": ("경기", "고양일산서구", "31104", "고양시 일산서구"),
    "310800": ("경기", "의정부시", "31030", "의정부시"),
    "312200": ("경기", "파주시", "31200", "파주시"),
    "312300": ("경기", "김포시", "31230", "김포시"),
    "312700": ("경기", "양주시", "31260", "양주시"),
    "110016": ("서울", "종로구", "11010", "종로구"),
    "110012": ("서울", "성북구", "11080", "성북구"),
    "110024": ("서울", "강북구", "11090", "강북구"),
    "110006": ("서울", "도봉구", "11100", "도봉구"),
    "110015": ("서울", "은평구", "11120", "은평구"),
    "110010": ("서울", "서대문구", "11130", "서대문구"),
    "110009": ("서울", "마포구", "11140", "마포구"),
    "110003": ("서울", "강서구", "11160", "강서구"),
}
BASE_COLUMNS = [ID, "요양기관명", "종별코드", "종별코드명", "시도코드명",
                "시군구코드", "시군구코드명", "주소", "전화번호", "좌표(X)", "좌표(Y)"]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def find_one(root, pattern):
    matches = list(root.rglob(pattern))
    require(len(matches) == 1, f"원자료가 정확히 1개여야 함: {pattern}, {len(matches)}개")
    return matches[0]


def clean(value):
    return "" if value is None else str(value).strip()


def read_workbook(path, columns, ids=None):
    """전국 자료를 순회하되 필요한 열·기관 행만 메모리에 보관한다."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        positions = [header.index(column) for column in columns]
        id_position = header.index(ID)
        records, total = [], 0
        for row in tqdm(rows, total=sheet.max_row - 1, desc=path.name[:16], mininterval=2):
            if all(value is None for value in row):
                continue
            total += 1
            if ids is None or clean(row[id_position]) in ids:
                records.append([clean(row[position]) for position in positions])
        return pd.DataFrame(records, columns=columns), total
    finally:
        workbook.close()


def read_boundaries(path):
    raw = gpd.read_file(path, engine="pyogrio")
    codes = {value[2]: value[3] for value in REGIONS.values()}
    require(raw.crs is not None, "경계 좌표계 누락")
    selected = raw.loc[raw.SIGUNGU_CD.astype(str).isin(codes)].copy()
    require(len(selected) == 15 and selected.SIGUNGU_CD.is_unique, "15개 시군구 경계 누락/중복")
    require(set(selected.BASE_DATE.astype(str)) == {"20250630"}, "경계 기준일 불일치")
    for row in selected.itertuples():
        require(row.SIGUNGU_NM == codes[str(row.SIGUNGU_CD)], "SGIS 코드·지역명 불일치")
    require(selected.geometry.notna().all() and (~selected.geometry.is_empty).all(), "빈 경계")
    selected = selected.to_crs(5179)
    selected.geometry = selected.geometry.map(lambda g: shapely.make_valid(g) if not g.is_valid else g)
    require(selected.is_valid.all() and (selected.area > 0).all(), "경계 도형 오류")
    return selected.set_index("SIGUNGU_CD"), shapely.union_all(selected.geometry.values)


def coordinate_valid(longitude, latitude):
    # 대한민국 원자료의 좌표 순서·결측·0·범위 오류를 함께 검출한다.
    return longitude.between(124, 132) & latitude.between(33, 40)


class KakaoGeocoder:
    def __init__(self, cache_path, offline):
        self.path, self.offline = cache_path, offline
        self.cache = {"schema_version": 1, "endpoint": ENDPOINT, "queries": {}}
        if cache_path.exists():
            self.cache = json.loads(cache_path.read_text(encoding="utf-8"))
        require(self.cache.get("endpoint") == ENDPOINT and self.cache.get("schema_version") == 1,
                "카카오 응답 원자료 형식 불일치")
        self.session = requests.Session()
        retry = Retry(total=2, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504],
                      allowed_methods=["GET"])
        self.session.mount("https://", HTTPAdapter(max_retries=retry))
        self.api_calls, self.cache_hits, self.records = 0, 0, []

    def query(self, address):
        if address in self.cache["queries"]:
            self.cache_hits += 1
            return self.cache["queries"][address]
        require(not self.offline, f"저장된 카카오 응답이 없음: {address}")
        key = os.getenv("KAKAO_API_KEY", "").strip()
        require(bool(key), "프로젝트 루트 .env 또는 환경변수에 KAKAO_API_KEY 설정 필요")
        response = self.session.get(ENDPOINT, headers={"Authorization": f"KakaoAK {key}"},
                                    params={"query": address}, timeout=(10, 30), allow_redirects=False)
        self.api_calls += 1
        # 인증 실패를 '주소 검색 결과 없음'으로 숨기지 않는다. 키·헤더는 출력하지 않는다.
        require(response.status_code == 200, f"카카오 주소 API 오류: HTTP {response.status_code}")
        body = response.json()
        require(isinstance(body.get("documents"), list), "카카오 응답 documents 형식 오류")
        entry = {"retrieved_at_utc": datetime.now(timezone.utc).isoformat(), "response": body}
        self.cache["queries"][address] = entry
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.path.with_suffix(".json.tmp")
        temporary.write_text(json.dumps(self.cache, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        temporary.replace(self.path)
        return entry

    def get_lat_lng(self, row):
        address = row["주소"].strip()
        require(bool(address), f"지오코딩할 주소가 없음: {row['요양기관명']}")
        # 원주소가 검색되지 않을 때만 층·호·괄호 설명을 제외한 주소로 재검색한다.
        simplified = re.sub(r"\([^)]*\)", "", address).split(",", 1)[0].strip()
        queries = list(dict.fromkeys([address, simplified]))
        attempted = []
        for query in queries:
            entry = self.query(query)
            documents = entry["response"]["documents"]
            attempted.append(query)
            if not documents:
                continue
            match = documents[0]  # 사용자 지정: 주소 검색의 첫 번째 결과
            require(match.get("address_type") in {"ROAD_ADDR", "REGION_ADDR"},
                    f"건물/지번 주소가 아닌 검색 결과: {row['요양기관명']}")
            require(len({(d.get('x'), d.get('y')) for d in documents}) == 1,
                    f"서로 다른 위치의 주소 검색 결과가 여러 개임: {row['요양기관명']}")
            longitude, latitude = float(match["x"]), float(match["y"])
            require(124 <= longitude <= 132 and 33 <= latitude <= 40, "카카오 반환 좌표 범위 오류")
            self.records.append({"hospital_id": row[ID], "name": row["요양기관명"],
                                 "address": address, "queries": attempted,
                                 "matched_address": match.get("address_name", ""),
                                 "longitude": longitude, "latitude": latitude,
                                 "retrieved_at_utc": entry["retrieved_at_utc"]})
            return pd.Series([longitude, latitude], index=["longitude", "latitude"])
        raise ValueError(f"카카오 주소 검색 결과 없음: {row['요양기관명']} / {address}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=ROOT/'data')
    parser.add_argument("--offline", action="store_true", help="저장된 API 응답만 사용; 네트워크 요청 금지")
    args = parser.parse_args()
    data = data_path(args.data_dir)
    load_dotenv(internal(ROOT / ".env"), override=False)
    tqdm.pandas(desc="누락 좌표 카카오 지오코딩", mininterval=1)
    base_path = find_one(data / "raw/병원", "1.병원정보서비스 2025.3.xlsx")
    specialty_path = find_one(data / "raw/병원", "5.의료기관별상세정보서비스_03_진료과목정보 2025.3.xlsx")
    boundary_path = find_one(data / "raw/경계", "bnd_sigungu_00_2025_2Q.shp")
    inputs = [base_path, specialty_path] + sorted(p for p in boundary_path.parent.glob(boundary_path.stem + ".*")
                                                    if p.suffix.lower() in {".shp", ".shx", ".dbf", ".prj", ".cpg"})
    hashes = {p.relative_to(data).as_posix(): sha256(p) for p in inputs}
    boundaries, study_area = read_boundaries(boundary_path)

    base, national_count = read_workbook(base_path, BASE_COLUMNS)
    require(base[ID].ne("").all() and base[ID].is_unique, "병원 기관 ID 결측/중복")
    regional = base.loc[base["시군구코드"].isin(REGIONS)].copy()
    for code, group in regional.groupby("시군구코드"):
        expected = REGIONS[code]
        require(set(group["시도코드명"]) == {expected[0]} and set(group["시군구코드명"]) == {expected[1]},
                f"HIRA 코드와 지역명이 맞지 않음: {code}")
    require(regional["요양기관명"].ne("").all(), "기관명 결측")
    specialties, specialty_total = read_workbook(
        specialty_path, [ID, "진료과목코드", "진료과목코드명", "과목별 전문의수"], set(regional[ID]))
    specialty_coverage = set(specialties[ID])
    specialties["진료과목코드"] = specialties["진료과목코드"].str.zfill(2)
    selected = specialties.loc[specialties["진료과목코드"].isin(DEPARTMENTS)].copy()
    require(not selected.duplicated([ID, "진료과목코드"]).any(), "같은 기관·진료과의 중복 행")
    require(selected["진료과목코드명"].eq(selected["진료과목코드"].map(DEPARTMENTS)).all(), "진료과 코드·이름 불일치")
    selected["count"] = pd.to_numeric(selected["과목별 전문의수"], errors="coerce")
    require(selected["count"].notna().all() and (selected["count"] >= 0).all()
            and (selected["count"] % 1 == 0).all(), "대상 진료과 전문의 수 결측/음수/비정수")
    pivot = selected.pivot(index=ID, columns="진료과목코드", values="count").reindex(columns=["01", "05"])
    pivot.columns = ["internal_medicine_specialists", "orthopedic_specialists"]
    joined = regional.merge(pivot, left_on=ID, right_index=True, how="left", validate="one_to_one")
    eligible = joined.loc[(joined.internal_medicine_specialists >= 1) | (joined.orthopedic_specialists >= 1)].copy()
    compact_names = eligible["요양기관명"].str.replace(r"\s+", "", regex=True)
    name_excluded = compact_names.str.contains("치과|요양병원", regex=True)
    excluded = eligible.loc[name_excluded]
    candidates = eligible.loc[~name_excluded].copy().reset_index(drop=True)
    candidates["longitude"] = pd.to_numeric(candidates["좌표(X)"], errors="coerce")
    candidates["latitude"] = pd.to_numeric(candidates["좌표(Y)"], errors="coerce")
    invalid = ~coordinate_valid(candidates.longitude, candidates.latitude)
    geocoder = KakaoGeocoder(data / "raw/병원/kakao_address_responses.json", args.offline)
    if invalid.any():
        candidates.loc[invalid, ["longitude", "latitude"]] = candidates.loc[invalid].progress_apply(
            geocoder.get_lat_lng, axis=1)
    require(coordinate_valid(candidates.longitude, candidates.latitude).all(), "미해결 좌표 존재")
    points = gpd.GeoSeries(gpd.points_from_xy(candidates.longitude, candidates.latitude), crs=4326).to_crs(5179)
    inside = points.map(study_area.covers)
    outside_records = candidates.loc[~inside, [ID, "요양기관명", "주소", "longitude", "latitude"]].to_dict("records")
    # 좌표를 경계에 스냅하지 않는다. 경계 밖은 결과에서 제외하고 표준출력으로 확인한다.
    candidates["x_5179"], candidates["y_5179"] = points.x, points.y
    mismatches = []
    for index, row in candidates.iterrows():
        code = REGIONS[row["시군구코드"]][2]
        if inside.iloc[index] and not boundaries.loc[code].geometry.covers(points.iloc[index]):
            mismatches.append({"hospital_id": row[ID], "name": row["요양기관명"], "address": row["주소"]})
    final = candidates.loc[inside].copy()
    final["is_goyang"] = final["시군구코드"].isin(["311901", "311902", "311903"]).astype(int)
    final["reference_month"] = "2025-03"
    final = final.rename(columns={ID: "hospital_id", "요양기관명": "name", "종별코드": "hospital_type_code",
                                  "종별코드명": "hospital_type", "시도코드명": "sido_name",
                                  "시군구코드": "hira_sigungu_code", "시군구코드명": "sigungu_name",
                                  "주소": "address", "전화번호": "telephone"})
    for column in pivot.columns:
        final[column] = final[column].astype("Int64")  # 없는 진료과 행을 임의로 0으로 바꾸지 않음.
    columns = ["hospital_id", "name", "hospital_type_code", "hospital_type", "sido_name", "hira_sigungu_code",
               "sigungu_name", "is_goyang", "address", "telephone", "internal_medicine_specialists",
               "orthopedic_specialists", "longitude", "latitude", "x_5179", "y_5179", "reference_month"]
    final = final[columns].sort_values(["hira_sigungu_code", "name", "hospital_id"]).reset_index(drop=True)
    require(final.hospital_id.is_unique and len(final) > 0, "최종 기관 ID 중복/빈 결과")
    require(not final.name.str.replace(r"\s+", "", regex=True).str.contains("치과|요양병원").any(), "제외 기관명 잔존")
    require(((final.internal_medicine_specialists >= 1) | (final.orthopedic_specialists >= 1)).fillna(False).all(),
            "최종 전문의 조건 오류")
    override_log = validate_geocodes(final)
    output = data / "processed/medical_institutions.csv"
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".csv.tmp")
    final.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
    readback = pd.read_csv(temporary, dtype=str, keep_default_na=False)
    require(list(readback.columns) == columns and len(readback) == len(final), "CSV 저장 구조 불일치")
    require(readback.hospital_id.tolist() == final.hospital_id.tolist(), "CSV ID 저장 불일치")
    for column in ["longitude", "latitude", "x_5179", "y_5179"]:
        require((pd.to_numeric(readback[column]) - final[column]).abs().max() < 1e-8, "CSV 좌표 정밀도 손실")
    require(all(sha256(p) == hashes[p.relative_to(data).as_posix()] for p in inputs), "입력 원자료 변경")
    temporary.replace(output)
    report = {
        "run_at_utc": datetime.now(timezone.utc).isoformat(), "input_sha256": hashes,
        "code_sha256": sha256(Path(__file__)), "output_sha256": sha256(output),
        "counts": {"national_hospitals": national_count, "national_specialty_rows": specialty_total,
                   "regional_hospitals": len(regional), "regional_no_specialty_records": len(set(regional[ID]) - specialty_coverage),
                   "specialist_eligible": len(eligible), "name_excluded": len(excluded),
                   "before_spatial_filter": len(candidates), "coordinates_geocoded": int(invalid.sum()),
                   "outside_study_area": int((~inside).sum()), "within_area_region_mismatch": len(mismatches),
                   "final": len(final), "goyang": int(final.is_goyang.sum()),
                   "neighbors": int((final.is_goyang == 0).sum()),
                   "internal_medicine": int((final.internal_medicine_specialists >= 1).sum()),
                   "orthopedic": int((final.orthopedic_specialists >= 1).sum()),
                   "both_departments": int(((final.internal_medicine_specialists >= 1) & (final.orthopedic_specialists >= 1)).sum()),
                   "same_coordinate_rows": int(final.duplicated(["longitude", "latitude"], keep=False).sum())},
        "by_type": final.hospital_type.value_counts().to_dict(),
        "by_region": final.groupby(["sido_name", "sigungu_name"]).size().reset_index(name="count").to_dict("records"),
        "excluded_by_type": excluded["종별코드명"].value_counts().to_dict(),
        "excluded_by_name": excluded[[ID, "요양기관명", "종별코드명"]].to_dict("records"),
        "geocoding": geocoder.records, "api_calls": geocoder.api_calls, "cache_hits": geocoder.cache_hits,
        "cache_sha256": sha256(geocoder.path) if geocoder.path.exists() else None,
        "outside_study_area": outside_records, "region_mismatches": mismatches,
        "missing_specialist_cells": {c: int(final[c].isna().sum()) for c in pivot.columns},
        "checks": {"unique_ids": True, "specialist_filter": True, "name_filter": True,
                   "coordinate_complete": True, "study_area": True, "csv_readback": True,
                   "input_unchanged": True},
        "environment": {"python": sys.version.split()[0], "pandas": pd.__version__,
                        "geopandas": gpd.__version__, "shapely": shapely.__version__},
    }
    report['override_application_log'] = override_log
    save_report('medical_institutions', report)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
