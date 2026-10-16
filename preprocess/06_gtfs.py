# -*- coding: utf-8 -*-
"""KTDB 2025-03 GTFS와 raw SGIS 경계로 분석용 gtfs.sqlite를 생성한다.

실행: python "preprocess/06_gtfs.py"
입력은 data/raw만 사용. MD·이미지 생성이나 OSM 접속/가상 환승선 생성은 하지 않는다.
원본 ID를 보존하며 stop_times의 정수 키는 trips·stops를 통해 원본 ID와 연결된다.
GTX·고속철도를 제외하고 최대 2회 환승 및 80분 시간 하한으로 노선 후보를 추린다.
검사 실패 시 기존 결과를 덮어쓰지 않는다. 결과 요약은 JSON 표준출력.
"""
from __future__ import annotations

import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.preprocessing_common import (ROOT, initialize, internal, data_path, save_report,
    building_constants, validate_buildings, building_application_log, walk_constants, validate_walk)
initialize()
import argparse
from datetime import datetime, timezone
import hashlib
import importlib.metadata
from itertools import groupby
import json
from pathlib import Path
import sqlite3
import sys

import geopandas as gpd
import numpy as np
from openpyxl import load_workbook
import pandas as pd
from scipy.spatial import cKDTree
from scipy.sparse import csr_matrix
from scipy.sparse.csgraph import dijkstra
import shapely

REGIONS = {'31101': '고양시 덕양구', '31103': '고양시 일산동구', '31104': '고양시 일산서구',
    '31030': '의정부시', '31200': '파주시', '31230': '김포시', '31260': '양주시',
    '11010': '종로구', '11080': '성북구', '11090': '강북구', '11100': '도봉구',
    '11120': '은평구', '11130': '서대문구', '11140': '마포구', '11160': '강서구'}
# KTDB 설명서 4쪽. 국제 표준 코드와 원본 코드를 구분한다.
MODES = {0: ('local_bus', 3), 1: ('urban_rail', 1), 2: ('ferry', 4),
         3: ('intercity_bus', 3), 4: ('conventional_rail', 2),
         5: ('airport_bus', 3), 6: ('high_speed_rail', 2), 7: ('air', None)}
GTX_IDS = {'RR_ACC1_S-1-XA-1D', 'RR_ACC1_S-1-XA-1U',
           'RR_ACC1_S-1-XA-2D', 'RR_ACC1_S-1-XA-2U'}
MAX_TRANSFERS = 2
TIME_COLUMNS = ['trip_id', 'arrival_time', 'departure_time', 'stop_id',
                'stop_sequence', 'pickup_type', 'drop_off_type', 'timepoint']


def require(condition, message):
    if not condition:
        raise ValueError(message)


def log(message):
    print(message, file=sys.stderr, flush=True)


def sha256(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024*1024), b''):
            h.update(block)
    return h.hexdigest()


def one(root, pattern):
    found = list(root.rglob(pattern))
    require(len(found) == 1, f'{pattern}: 원본 1개 필요, 발견 {len(found)}')
    return found[0]


def read_txt(path):
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding='utf-8-sig')


def unique(frame, column):
    require(frame[column].ne('').all() and frame[column].is_unique, f'{column}: 빈 ID/중복 ID')


def read_details(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    def sheet_table(sheet, first):
        header, records = None, []
        for row in workbook[sheet].iter_rows(values_only=True):
            if row[0] == first:
                require(header is None, '노선 세부정보 헤더 중복')
                header = list(row)
            elif header and row[0] is not None:
                records.append(list(row))
        require(header is not None, f'{sheet}: 영문 헤더 없음')
        return pd.DataFrame(records, columns=header).fillna('').astype(str)
    try:
        details = sheet_table('유형', 'route_id')
        categories = sheet_table('route_cate정의', 'RT_CA')
    finally:
        workbook.close()
    duplicate_count = int(categories.duplicated().sum())
    categories = categories.drop_duplicates().reset_index(drop=True)
    categories.attrs['exact_duplicate_rows_removed'] = duplicate_count
    unique(details, 'route_id'); unique(categories, 'RT_CA')
    return details, categories


def source_tables(root, details_path):
    tables = {name: read_txt(root/f'{name}.txt') for name in ['agency', 'calendar', 'routes', 'trips', 'stops', 'transfers']}
    for name, key in [('agency', 'agency_id'), ('calendar', 'service_id'), ('routes', 'route_id'), ('trips', 'trip_id'), ('stops', 'stop_id')]:
        unique(tables[name], key)
    routes = tables['routes'].sort_values('route_id').reset_index(drop=True)
    trips = tables['trips'].sort_values('trip_id').reset_index(drop=True)
    stops = tables['stops'].sort_values('stop_id').reset_index(drop=True)
    routes['route_type'] = routes.route_type.astype(int)
    require(set(routes.route_type) <= set(MODES), '알 수 없는 KTDB 수단 코드')
    details, categories = read_details(details_path)
    details['route_type'] = details.route_type.astype(int)
    require(set(details.route_id) == set(routes.route_id), 'GTFS·노선 세부정보 ID 불일치')
    joined = routes.merge(details, on='route_id', suffixes=('', '_detail'), validate='one_to_one')
    for col in ['agency_id', 'route_short_name', 'route_long_name', 'route_type']:
        require(joined[col].equals(joined[col+'_detail']), f'노선 세부정보 {col} 불일치')
    detected = set(joined.loc[joined.route_short_name.str.contains('GTX', case=False) |
        joined.route_long_name.str.contains('GTX', case=False), 'route_id'])
    require(detected == GTX_IDS, 'GTX 노선 ID 변경: 노선 세부정보 재검토 필요')
    require(set(joined.loc[joined.route_id.isin(GTX_IDS), 'region_name']) == {'수도권'}, 'GTX 권역 불일치')
    joined['route_key'] = np.arange(1, len(joined)+1)
    joined['mode'] = joined.route_type.map(lambda n: MODES[n][0])
    joined['route_type_standard'] = joined.route_type.map(lambda n: MODES[n][1])
    joined['category_name'] = joined.route_cate.map(categories.set_index('RT_CA').RT_CA_NM).fillna('')
    # 원본 사전에 일부 지역의 버스 세부 코드가 없다. 수단 코드는 별도로
    # 확인되므로 노선을 버리지 않고 원본 세부 코드와 빈 분류명을 보존한다.
    joined['mode_eligible'] = (~joined.route_type.isin([2, 3, 5, 6, 7]) & ~joined.route_id.isin(GTX_IDS)).astype(int)
    joined['is_gtx'] = joined.route_id.isin(GTX_IDS).astype(int)
    trips['trip_key'] = np.arange(1, len(trips)+1)
    trips['route_key'] = trips.route_id.map(joined.set_index('route_id').route_key)
    require(trips.route_key.notna().all(), '운행편의 노선 참조 누락')
    require(trips.service_id.isin(tables['calendar'].service_id).all(), '운행편의 서비스 참조 누락')
    require(joined.agency_id.isin(tables['agency'].agency_id).all(), '노선의 기관 참조 누락')
    require(len(tables['calendar']) == 1 and tables['calendar'].iloc[0].to_dict() == {
        'service_id': 'B1', **{d: '1' for d in ['monday','tuesday','wednesday','thursday','friday','saturday','sunday']},
        'start_date': '20170101', 'end_date': '20301231'}, '달력 원본 변경: 평일 표본 해석 재검토 필요')
    require(set(tables['agency'].agency_timezone) == {'Asia/Seoul'}, '시간대 불일치')
    stops['stop_key'] = np.arange(1, len(stops)+1)
    stops[['stop_lat', 'stop_lon']] = stops[['stop_lat', 'stop_lon']].astype(float)
    require(np.isfinite(stops[['stop_lat','stop_lon']]).all().all() and
            stops.stop_lat.between(-90,90).all() and stops.stop_lon.between(-180,180).all(), '정류장 좌표 오류')
    tables.update(routes=joined, trips=trips, stops=stops, route_categories=categories)
    return tables


def route_stop_pairs(path, trips, stops, chunksize):
    """정차 관계와 같은 운행편의 연속 정류장 사이 최소 주행시간을 구한다."""
    trip_routes = trips.set_index('trip_id').route_key
    stop_keys = stops.set_index('stop_id').stop_key
    pairs, rows, edge_parts, previous = set(), 0, [], None
    for x in pd.read_csv(path, usecols=['trip_id','stop_id','arrival_time','departure_time'], dtype=str, keep_default_na=False, chunksize=chunksize):
        a, b = x.trip_id.map(trip_routes), x.stop_id.map(stop_keys)
        require(a.notna().all() and b.notna().all(), '시간표의 운행편/정류장 참조 누락')
        pairs.update(zip(a, b))
        x['route_key'], x['stop_key'] = a, b
        x['arrival_s'], x['departure_s'] = seconds(x.arrival_time), seconds(x.departure_time)
        y = pd.concat([previous,x],ignore_index=True) if previous is not None else x
        same = y.trip_id.eq(y.trip_id.shift())
        edges = pd.DataFrame({'route_key':y.route_key,'from_stop':y.stop_key.shift(),
            'to_stop':y.stop_key,'travel_s':y.arrival_s-y.departure_s.shift()}).loc[same].astype('int64')
        require(edges.travel_s.ge(0).all(), '후보 추출용 정차 간 시간 역전')
        edge_parts.append(edges.groupby(['route_key','from_stop','to_stop'],as_index=False).travel_s.min())
        previous = x.tail(1)
        rows += len(x)
        if rows % (chunksize*8) == 0:
            log(f'공간 후보용 시간표 {rows:,}행 확인')
    edges = pd.concat(edge_parts,ignore_index=True).groupby(['route_key','from_stop','to_stop'],as_index=False).travel_s.min()
    return pd.DataFrame(sorted(pairs), columns=['route_key','stop_key']), edges, rows


def ride_lower_bounds(pair_stops, edge_u, edge_v, edge_cost, access, limit):
    """노선별 정류장 그래프에서 한 번 탑승했을 때의 낙관적 도착시간."""
    n = len(pair_stops)
    starts = np.flatnonzero(access[pair_stops] <= limit)
    graph = csr_matrix((np.r_[edge_cost,access[pair_stops[starts]]],
        (np.r_[edge_u,np.full(len(starts),n)],np.r_[edge_v,starts])),shape=(n+1,n+1))
    # 희소행렬의 명시적 0 비용 간선도 보존: 고양시 안에서는 접근 하한이 0.
    return dijkstra(graph,directed=True,indices=n,limit=limit)[:n]


def select_routes(tables, pairs, edges, boundary, radius, max_boardings, horizon_s, walk_speed):
    b = gpd.read_file(boundary).to_crs(5179)
    selected = b[b.SIGUNGU_CD.astype(str).isin(REGIONS)]
    require(len(selected) == 15 and selected.SIGUNGU_CD.is_unique, '경계 15개 누락/중복')
    require(set(selected.BASE_DATE.astype(str)) == {'20250630'}, '경계 기준일 불일치')
    for row in selected.itertuples():
        require(REGIONS[str(row.SIGUNGU_CD)] == row.SIGUNGU_NM and row.geometry.is_valid, '경계 코드/도형 오류')
    core = shapely.union_all(selected.geometry.values)
    goyang = shapely.union_all(selected[selected.SIGUNGU_CD.astype(str).str.startswith('3110')].geometry.values)
    stops, routes = tables['stops'], tables['routes']
    points = gpd.GeoSeries(gpd.points_from_xy(stops.stop_lon, stops.stop_lat), crs=4326).to_crs(5179)
    xy = np.c_[points.x, points.y]
    stops['x_5179'], stops['y_5179'] = xy[:,0], xy[:,1]
    stops['in_goyang'] = shapely.covers(goyang, points.values).astype(int)
    stops['in_institution_region'] = shapely.covers(core, points.values).astype(int)
    origin_distance = shapely.distance(points.values,goyang)
    seed = origin_distance <= radius
    stops['origin_search_candidate'] = seed.astype(int)
    # 시간 하한은 대기·정차·승하차 제한을 완화한다. 실제 경로 비용으로 쓰지 않는다.
    eligible_routes = routes.loc[routes.mode_eligible == 1,'route_key']
    eligible_pairs = pairs[pairs.route_key.isin(eligible_routes)].reset_index(drop=True)
    edges = edges[edges.route_key.isin(eligible_routes)]
    stride = len(stops)+1
    codes = eligible_pairs.route_key.to_numpy()*stride + eligible_pairs.stop_key.to_numpy()
    u_codes = edges.route_key.to_numpy()*stride + edges.from_stop.to_numpy()
    v_codes = edges.route_key.to_numpy()*stride + edges.to_stop.to_numpy()
    u,v = np.searchsorted(codes,u_codes),np.searchsorted(codes,v_codes)
    require(np.array_equal(codes[u],u_codes) and np.array_equal(codes[v],v_codes),'노선 정차 그래프 참조 오류')
    pair_stops = eligible_pairs.stop_key.to_numpy()-1
    costs = edges.travel_s.to_numpy(dtype=float)
    stop_index = stops.set_index('stop_id').stop_key
    transfers = tables['transfers'].copy()
    for column in ['from_stop_id','to_stop_id']:
        require(transfers[column].isin(stop_index.index).all(), '환승 정류장 참조 누락')
    transfer_keys = [(int(stop_index[a])-1,int(stop_index[b])-1) for a,b in
                     zip(transfers.from_stop_id,transfers.to_stop_id)]
    def transfer_closure(elapsed):
        # 추출 단계에서는 방향/시간을 완화하여 과소 추출을 방지한다.
        # 결과 transfers 테이블은 원본 방향과 최소환승시간을 그대로 유지한다.
        elapsed = elapsed.copy()
        while True:
            changed = False
            for a,b in transfer_keys:
                if elapsed[a] != elapsed[b]:
                    elapsed[a] = elapsed[b] = min(elapsed[a],elapsed[b]); changed = True
            if not changed:
                return elapsed
    tree = cKDTree(xy)
    def walk_relaxation(elapsed):
        before = transfer_closure(elapsed)
        after = before.copy()
        origins = np.flatnonzero(before<=horizon_s)
        for start in range(0,len(origins),256):
            batch = origins[start:start+256]
            for i,neighbors in zip(batch,tree.query_ball_point(xy[batch],radius+1e-6)):
                neighbors = np.asarray(neighbors,dtype=int)
                lower = before[i]+np.linalg.norm(xy[neighbors]-xy[i],axis=1)/walk_speed
                after[neighbors] = np.minimum(after[neighbors],lower)
        after = transfer_closure(after)
        after[after>horizon_s] = np.inf
        return after
    stages, boarding = [], {}
    elapsed = np.where(seed,origin_distance/walk_speed,np.inf)
    for step in range(1, max_boardings+1):
        access = transfer_closure(elapsed) if step==1 else walk_relaxation(elapsed)
        lower = ride_lower_bounds(pair_stops,u,v,costs,access,horizon_s)
        reachable = np.isfinite(lower)
        found = set(eligible_pairs.loc[reachable,'route_key'])
        new = found - set(boarding)
        boarding.update({int(k): step for k in new})
        arrived = np.full(len(stops),np.inf)
        np.minimum.at(arrived,pair_stops,lower)
        elapsed = np.minimum(elapsed,arrived)
        require(found, '접근 가능한 노선 후보 없음')
        stages.append({'boarding_round': step,'transfer_round':step-1,'new_routes':len(new),
            'total_routes':len(boarding),'stops_within_time_lower_bound':int(np.isfinite(arrived).sum())})
        log(f'최대 {step-1}회 환승 / {horizon_s/60:g}분 하한 후보 노선 {len(boarding):,}개')
    routes['min_boarding_round'] = routes.route_key.map(boarding).fillna(0).astype(int)
    routes['selected'] = (routes.min_boarding_round > 0).astype(int)
    routes['selection_reason'] = np.select([routes.is_gtx==1, routes.route_type==6, routes.route_type==3, routes.route_type==5,
        routes.mode_eligible==0, routes.selected==0],
        ['gtx_excluded','high_speed_rail_excluded','intercity_express_bus_excluded','airport_bus_excluded','mode_outside_study','outside_time_transfer_envelope'], default='retained_complete_route')
    selected_pairs = pairs[pairs.route_key.isin(boarding)]
    stop_set = set(selected_pairs.stop_key)
    # OSM의 초기 추출 사각형을 raw 경계로 재생성. 이 표시가 실제 연결 여부를 의미하지 않는다.
    osm_bounds = gpd.GeoSeries([core.buffer(2000).envelope], crs=5179).to_crs(4326).total_bounds
    osm_inside = stops.stop_lon.between(osm_bounds[0],osm_bounds[2]) & stops.stop_lat.between(osm_bounds[1],osm_bounds[3])
    tables['stops'] = stops.loc[stops.stop_key.isin(stop_set)].copy()
    tables['trips'] = tables['trips'].loc[tables['trips'].route_key.isin(boarding)].copy()
    return selected_pairs, {'stages': stages, 'seed_stops': int(seed.sum()),
        'ride_lower_bound_edges':len(edges),'ride_lower_bound_states':len(eligible_pairs),
        'selected_stops_outside_institution_region': int((tables['stops'].in_institution_region==0).sum()),
        'selected_stops_outside_osm_initial_bbox': int((stops.stop_key.isin(stop_set)&~osm_inside).sum()),
        'goyang_boundary_wkt_5179': goyang.wkt, 'institution_boundary_wkt_5179': core.wkt}


def seconds(values):
    require(values.str.fullmatch(r'\d{2,3}:[0-5]\d:[0-5]\d').all(), '시간 형식 오류')
    parts = values.str.split(':', expand=True).astype('int64')
    return parts[0]*3600 + parts[1]*60 + parts[2]


SCHEMA = '''
CREATE TABLE agency(agency_id TEXT PRIMARY KEY,agency_name TEXT,agency_url TEXT,agency_timezone TEXT);
CREATE TABLE calendar(service_id TEXT PRIMARY KEY,monday INTEGER,tuesday INTEGER,wednesday INTEGER,thursday INTEGER,friday INTEGER,saturday INTEGER,sunday INTEGER,start_date TEXT,end_date TEXT);
CREATE TABLE routes(route_id TEXT PRIMARY KEY,agency_id TEXT REFERENCES agency(agency_id),route_short_name TEXT,route_long_name TEXT,route_type_raw INTEGER,route_type INTEGER,mode TEXT,region_name TEXT,route_cate TEXT,category_name TEXT,min_boarding_round INTEGER);
CREATE TABLE stops(stop_key INTEGER PRIMARY KEY,stop_id TEXT UNIQUE NOT NULL,stop_name TEXT,stop_lat REAL,stop_lon REAL,x_5179 REAL,y_5179 REAL,in_goyang INTEGER,in_institution_region INTEGER,origin_search_candidate INTEGER);
CREATE TABLE patterns(pattern_id TEXT PRIMARY KEY,route_id TEXT REFERENCES routes(route_id),stop_count INTEGER,first_stop_key INTEGER REFERENCES stops(stop_key),last_stop_key INTEGER REFERENCES stops(stop_key));
CREATE TABLE pattern_stops(pattern_id TEXT REFERENCES patterns(pattern_id),pattern_sequence INTEGER,stop_key INTEGER REFERENCES stops(stop_key),pickup_type INTEGER,drop_off_type INTEGER,PRIMARY KEY(pattern_id,pattern_sequence)) WITHOUT ROWID;
CREATE TABLE trips(trip_key INTEGER PRIMARY KEY,trip_id TEXT UNIQUE NOT NULL,route_id TEXT REFERENCES routes(route_id),service_id TEXT REFERENCES calendar(service_id),pattern_id TEXT REFERENCES patterns(pattern_id),stop_count INTEGER,first_departure_s INTEGER,last_arrival_s INTEGER);
CREATE TABLE stop_times(trip_key INTEGER REFERENCES trips(trip_key),stop_sequence INTEGER CHECK(stop_sequence>=0),stop_key INTEGER REFERENCES stops(stop_key),arrival_s INTEGER CHECK(arrival_s>=0),departure_s INTEGER CHECK(departure_s>=arrival_s),pickup_type INTEGER CHECK(pickup_type IN(0,1)),drop_off_type INTEGER CHECK(drop_off_type IN(0,1)),timepoint INTEGER CHECK(timepoint=1),PRIMARY KEY(trip_key,stop_sequence)) WITHOUT ROWID;
CREATE TABLE transfers(from_stop_key INTEGER REFERENCES stops(stop_key),to_stop_key INTEGER REFERENCES stops(stop_key),transfer_type INTEGER,min_transfer_time INTEGER CHECK(min_transfer_time>=0),PRIMARY KEY(from_stop_key,to_stop_key)) WITHOUT ROWID;
CREATE TABLE processing_metadata(key TEXT PRIMARY KEY,value_json TEXT);
CREATE TABLE qa_checks(check_name TEXT PRIMARY KEY,passed INTEGER,value_json TEXT);
'''


def insert_frame(c, table, frame):
    c.executemany(f'INSERT INTO {table} VALUES ({",".join("?" for _ in frame.columns)})', frame.itertuples(index=False,name=None))


def write_small_tables(c, tables):
    insert_frame(c,'agency',tables['agency'])
    cal = tables['calendar'].copy()
    cal.iloc[:,1:8] = cal.iloc[:,1:8].astype(int)
    insert_frame(c,'calendar',cal)
    routes = tables['routes']
    out = routes.loc[routes.selected==1, ['route_id','agency_id','route_short_name','route_long_name','route_type',
        'route_type_standard','mode','region_name','route_cate','category_name','min_boarding_round']].copy()
    out['route_type_standard'] = out.route_type_standard.astype(int)
    insert_frame(c,'routes',out)
    decisions = routes[['route_id','route_short_name','route_long_name','agency_id','route_type','route_type_standard',
        'mode','region_name','route_cate','category_name','is_gtx','selected','min_boarding_round','selection_reason']]
    decisions.to_sql('route_decisions',c,index=False)
    tables['route_categories'].to_sql('route_categories',c,index=False)
    insert_frame(c,'stops',tables['stops'][['stop_key','stop_id','stop_name','stop_lat','stop_lon','x_5179','y_5179',
        'in_goyang','in_institution_region','origin_search_candidate']])
    t = tables['trips'][['trip_key','trip_id','route_id','service_id']].copy()
    for col in ['pattern_id','stop_count','first_departure_s','last_arrival_s']:t[col] = None
    insert_frame(c,'trips',t)
    s = tables['stops'].set_index('stop_id').stop_key
    transfer = tables['transfers'].copy()
    require(~transfer.duplicated(['from_stop_id','to_stop_id']).any(), '중복 환승 쌍')
    transfer[['transfer_type','min_transfer_time']] = transfer[['transfer_type','min_transfer_time']].astype(int)
    require(set(transfer.transfer_type) == {2} and transfer.min_transfer_time.ge(0).all(), '환승 유형/시간 재검토 필요')
    transfer['selected'] = (transfer.from_stop_id.isin(s.index)&transfer.to_stop_id.isin(s.index)).astype(int)
    transfer['selection_reason'] = np.where(transfer.selected==1,'both_stops_retained','unused_stop_after_route_selection')
    transfer.to_sql('transfer_decisions',c,index=False)
    rows = [(int(s[a]),int(s[b]),int(t),int(m)) for a,b,t,m in
            transfer.loc[transfer.selected==1].iloc[:,:4].itertuples(index=False,name=None)]
    c.executemany('INSERT INTO transfers VALUES (?,?,?,?)',rows)
    return len(rows)


def load_times(c, path, all_trips, all_stops, selected_keys, chunksize):
    trip_map = all_trips.set_index('trip_id').trip_key
    stop_map = all_stops.set_index('stop_id').stop_key
    counts = np.zeros(len(all_trips)+1,dtype='int64')
    blocks = counts.copy()
    raw_rows, kept_rows, overnight, kept_overnight = 0,0,0,0
    previous = None
    digest = hashlib.sha256()
    columns = ['trip_key','stop_sequence','stop_key','arrival_s','departure_s','pickup_type','drop_off_type','timepoint']
    for x in pd.read_csv(path,dtype=str,keep_default_na=False,chunksize=chunksize):
        require(list(x.columns)==TIME_COLUMNS and not x.eq('').any().any(), '시간표 열/빈값 오류')
        x['trip_key'] = x.trip_id.map(trip_map)
        x['stop_key'] = x.stop_id.map(stop_map)
        require(x.trip_key.notna().all() and x.stop_key.notna().all(), '시간표 ID 참조 누락')
        for col in ['stop_sequence','pickup_type','drop_off_type','timepoint']:
            require(x[col].str.fullmatch(r'\d+').all(), f'{col}: 정수 형식 오류')
            x[col] = x[col].astype('int64')
        require(x.pickup_type.isin([0,1]).all() and x.drop_off_type.isin([0,1]).all() and x.timepoint.eq(1).all(), '승하차/시각 정확도 코드 변경')
        x['arrival_s'],x['departure_s'] = seconds(x.arrival_time),seconds(x.departure_time)
        require(x.arrival_s.le(x.departure_s).all(), '같은 정차의 도착/출발 시각 역전')
        y = pd.concat([previous,x],ignore_index=True) if previous is not None else x
        same = y.trip_key.eq(y.trip_key.shift())
        require((~same | y.stop_sequence.gt(y.stop_sequence.shift())).all(), '정차 순서 중복/역전: 원자료 검토 필요')
        require((~same | y.arrival_s.ge(y.departure_s.shift())).all(), '정차 간 시각 역전')
        starts = (~same).iloc[-len(x):].to_numpy()
        counts += np.bincount(x.trip_key.to_numpy(),minlength=len(counts))
        blocks += np.bincount(x.loc[starts,'trip_key'].to_numpy(),minlength=len(blocks))
        previous = x.tail(1)
        raw_rows += len(x); overnight += int(x.departure_s.ge(86400).sum())
        out = x.loc[x.trip_key.isin(selected_keys),columns].astype('int64')
        kept_rows += len(out); kept_overnight += int(out.departure_s.ge(86400).sum())
        # 원본 순서가 trip_key·stop_sequence 순인지는 별도 요구하지 않는다.
        # 이 해시는 입력에서 선별한 행의 순서대로 계산한 실행 기록이다.
        digest.update(out.to_numpy(dtype='<i8').tobytes())
        insert_frame(c,'stop_times',out)
        if raw_rows % (chunksize*4) == 0:log(f'시간표 검사 {raw_rows:,}행 / 저장 {kept_rows:,}행')
    require((counts[1:]>=2).all(), '정차 2개 미만/시간표 없는 운행편')
    require((blocks[1:]==1).all(), '분산된 운행편 블록: 정렬 후 재검토 필요')
    c.commit()
    return {'raw_rows':raw_rows,'retained_rows':kept_rows,'raw_overnight_rows':overnight,
            'retained_overnight_rows':kept_overnight,'source_order_loaded_sha256':digest.hexdigest()}, counts


def build_patterns(c):
    routes = dict(c.execute('SELECT trip_key,route_id FROM trips'))
    patterns, updates = {}, []
    rows = c.execute('SELECT trip_key,stop_sequence,stop_key,arrival_s,departure_s,pickup_type,drop_off_type FROM stop_times ORDER BY trip_key,stop_sequence')
    for key, items in groupby(rows,key=lambda x:x[0]):
        trip = list(items)
        sequence = [(int(v[2]),int(v[5]),int(v[6])) for v in trip]
        canonical = json.dumps([routes[key],sequence],separators=(',',':'))
        pattern = hashlib.sha256(canonical.encode()).hexdigest()
        if pattern not in patterns:
            patterns[pattern] = (routes[key],sequence)
        else:
            require(patterns[pattern] == (routes[key],sequence),'패턴 해시 충돌')
        updates.append((pattern,len(trip),trip[0][4],trip[-1][3],key))
    c.executemany('INSERT INTO patterns VALUES (?,?,?,?,?)',
        [(p,route,len(seq),seq[0][0],seq[-1][0]) for p,(route,seq) in patterns.items()])
    for p,(_,seq) in patterns.items():
        c.executemany('INSERT INTO pattern_stops VALUES (?,?,?,?,?)',[(p,i+1,*s) for i,s in enumerate(seq)])
    c.executemany('UPDATE trips SET pattern_id=?,stop_count=?,first_departure_s=?,last_arrival_s=? WHERE trip_key=?',updates)
    return len(patterns)


def validate(c, counts, report):
    checks = {}
    def check(name,condition,value):
        require(bool(condition),name)
        checks[name] = value
    check('sqlite_integrity',c.execute('PRAGMA integrity_check').fetchone()[0]=='ok','ok')
    foreign = c.execute('PRAGMA foreign_key_check').fetchall()
    check('foreign_keys',not foreign,len(foreign))
    actual = dict(c.execute('SELECT trip_key,count(*) FROM stop_times GROUP BY trip_key'))
    check('complete_trips',all(n==int(counts[k]) for k,n in actual.items()),len(actual))
    check('all_retained_trips_have_times',len(actual)==c.execute('SELECT count(*) FROM trips').fetchone()[0],len(actual))
    check('stop_times_row_count',sum(actual.values())==report['retained_rows'],sum(actual.values()))
    bad = c.execute('''SELECT count(*) FROM (SELECT trip_key,stop_sequence,arrival_s,
        lag(departure_s) OVER(PARTITION BY trip_key ORDER BY stop_sequence) prev FROM stop_times) WHERE arrival_s<prev''').fetchone()[0]
    check('saved_time_monotonicity',bad==0,bad)
    bad = c.execute('''SELECT count(*) FROM trips t JOIN patterns p USING(pattern_id)
        WHERE t.route_id<>p.route_id OR t.stop_count<>p.stop_count OR t.pattern_id IS NULL''').fetchone()[0]
    check('trip_patterns',bad==0 and c.execute('SELECT count(*) FROM trips WHERE pattern_id IS NULL').fetchone()[0]==0,bad)
    bad = c.execute('SELECT count(*) FROM routes r JOIN route_decisions d USING(route_id) WHERE d.is_gtx=1 OR d.selected<>1 OR r.route_type_raw IN(2,3,5,6,7) OR r.route_type NOT IN(1,2,3)').fetchone()[0]
    check('gtx_and_modes_excluded',bad==0,bad)
    bad = c.execute('SELECT count(*) FROM routes WHERE min_boarding_round NOT BETWEEN 1 AND ?',
                    (MAX_TRANSFERS+1,)).fetchone()[0]
    check('maximum_two_transfer_candidates',bad==0,bad)
    unused = c.execute('SELECT count(*) FROM stops s WHERE NOT EXISTS(SELECT 1 FROM stop_times t WHERE t.stop_key=s.stop_key)').fetchone()[0]
    check('no_unused_stops',unused==0,unused)
    # 정수 키를 원본 ID로 되돌려 조회할 수 있는 뷰를 제공한다.
    c.executescript('''
    CREATE VIEW stop_times_named AS SELECT t.trip_id,s.stop_id,x.stop_sequence,x.arrival_s,x.departure_s,
        x.pickup_type,x.drop_off_type,x.timepoint FROM stop_times x JOIN trips t USING(trip_key) JOIN stops s USING(stop_key);
    CREATE VIEW transfers_named AS SELECT a.stop_id from_stop_id,b.stop_id to_stop_id,t.transfer_type,t.min_transfer_time
        FROM transfers t JOIN stops a ON a.stop_key=t.from_stop_key JOIN stops b ON b.stop_key=t.to_stop_key;
    ''')
    for name,value in checks.items():
        c.execute('INSERT INTO qa_checks VALUES (?,?,?)',(name,1,json.dumps(value,ensure_ascii=False)))
    return checks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data-dir',type=Path,default=ROOT/'data')
    parser.add_argument('--candidate-radius-m',type=float,default=2000)
    parser.add_argument('--candidate-max-minutes',type=float,default=80)
    parser.add_argument('--candidate-walk-speed',type=float,default=1.2)
    parser.add_argument('--chunksize',type=int,default=500000)
    args = parser.parse_args()
    require(args.candidate_radius_m>=480 and args.candidate_max_minutes>=80 and
            args.candidate_walk_speed>=0.8 and args.chunksize>0,'추출 설정이 기본 연구 조건보다 작음')
    max_boardings = MAX_TRANSFERS + 1  # 첫 탑승 + 최대 2회 환승 = 최대 3회 탑승
    data = data_path(args.data_dir);raw = data/'raw'
    source = one(raw/'GTFS','stop_times.txt').parent
    details = one(raw/'GTFS','202503_GTFS_route세부정보.xlsx')
    manual = one(raw/'GTFS','202503_GTFS_설명서.pdf')
    boundary = one(raw/'경계','bnd_sigungu_00_2025_2Q.shp')
    inputs = sorted(source.glob('*.txt'))+[details,manual]+[boundary.with_suffix(s) for s in ['.shp','.shx','.dbf','.prj','.cpg']]
    hashes = {p.relative_to(data).as_posix():sha256(p) for p in inputs}
    log('GTFS·노선 세부정보·경계 읽기')
    tables = source_tables(source,details)
    raw_counts = {k:len(v) for k,v in tables.items() if k != 'route_categories'}
    all_trips,all_stops = tables['trips'].copy(),tables['stops'].copy()
    pairs,ride_edges,raw_time_rows = route_stop_pairs(source/'stop_times.txt',all_trips,all_stops,args.chunksize)
    pairs,spatial = select_routes(tables,pairs,ride_edges,boundary,args.candidate_radius_m,
        max_boardings,args.candidate_max_minutes*60,args.candidate_walk_speed)
    output = data/'processed/gtfs.sqlite';output.parent.mkdir(parents=True,exist_ok=True)
    temporary = output.with_name('gtfs.writing.sqlite')
    require(not temporary.exists(),'이전 임시 GTFS 출력 존재: 확인 후 정리 필요')
    c = sqlite3.connect(temporary)
    try:
        c.executescript('PRAGMA journal_mode=DELETE; PRAGMA synchronous=NORMAL; PRAGMA cache_size=-131072; PRAGMA temp_store=MEMORY;'+SCHEMA)
        write_small_tables(c,tables)
        selected_keys = set(tables['trips'].trip_key)
        report,source_trip_counts = load_times(c,source/'stop_times.txt',all_trips,all_stops,selected_keys,args.chunksize)
        require(report['raw_rows']==raw_time_rows,'두 차례 원자료 읽기 행 수 불일치')
        log('정차 패턴 생성 및 검색 인덱스 구성')
        patterns = build_patterns(c)
        c.executescript('''CREATE INDEX stop_times_stop_departure ON stop_times(stop_key,departure_s,trip_key);
            CREATE INDEX trips_route ON trips(route_id); CREATE INDEX trips_pattern ON trips(pattern_id);
            CREATE INDEX stops_name ON stops(stop_name); CREATE INDEX patterns_route ON patterns(route_id);''')
        c.commit()
        log('저장 시간표·참조·완전 운행편 검증')
        checks = validate(c,source_trip_counts,report)
        require(all(sha256(p)==hashes[p.relative_to(data).as_posix()] for p in inputs),'처리 중 원자료 변경')
        counts = {table:c.execute(f'SELECT count(*) FROM {table}').fetchone()[0] for table in
            ['agency','calendar','routes','stops','trips','stop_times','patterns','pattern_stops','transfers','route_decisions','transfer_decisions','route_categories']}
        metadata = {'created_at_utc':datetime.now(timezone.utc).isoformat(),'raw_counts':dict(raw_counts,stop_times=raw_time_rows),
            'counts':counts,'time_statistics':report,'spatial_selection':spatial,'checks':checks,
            'category_exact_duplicate_rows_removed':tables['route_categories'].attrs['exact_duplicate_rows_removed'],
            'unmapped_bus_categories':tables['routes'].loc[
                tables['routes'].route_type.isin([0,5]) & tables['routes'].category_name.eq('')
            ].groupby(['route_type','region_name','route_cate','selected']).size().rename('routes').reset_index().to_dict('records'),
            'parameters':{'candidate_radius_m':args.candidate_radius_m,'max_transfers':MAX_TRANSFERS,'max_boardings':max_boardings,
                'candidate_horizon_s':args.candidate_max_minutes*60,'candidate_walk_speed_mps':args.candidate_walk_speed,
                'candidate_scope_check': {'max_scenario_speed_mps':0.8,'continuous_walk_limit_s':600,'required_walk_distance_m':480,'max_diagnostic_horizon_s':4800,'total_walk_limit_s':None,'passed':True}, 'time_filter':'none_full_service_day','time_unit':'seconds_since_service_day_midnight','timezone':'Asia/Seoul'},
            'service_interpretation':'2025년 3월 평일 1일 표본. calendar의 전 요일/2017~2030은 분석 편의 코딩이며 실제 전 기간·주말 운행을 의미하지 않음.',
            'connection_status':'원본 철도 환승만 보존. 후보용 시간 하한은 실제 경로 비용이 아니며 버스 환승·역 내부·OSM 접속은 후속 처리.',
            'agency_interpretation':'A1/KTDB는 통합 제공기관 값. 실제 운수회사 식별 필드가 아님.',
            'spatial_scope_rule':'거주지는 고양시, 기관 선정 범위는 고양시+지정 12개 지역. GTX·고속철도·시외/고속버스·공항버스 제외. 노선별 연속 정차의 최소 주행시간과 직선 보행시간으로 계산한 하한이 80분 이내인 후보를 최대 2회 환승까지 확장. 대기·정차·승하차 제한 및 보행 장벽을 완화하며 원본 환승 연결도 추출 시에만 양방향 0초로 완화. 실제 경로가 아닌 낙관적 후보 검사. 선택한 노선의 전 운행편·전 정차·하루 전체 시각 보존. 시간예산·보행 조건·환승 상한·노선 개선안 변경 시 후보 범위 재검증.',
            'gtx_excluded':sorted(GTX_IDS),
            'high_speed_rail_excluded':sorted(tables['routes'].loc[tables['routes'].route_type==6,'route_id']),
            'route_type_mapping':MODES,'input_sha256':hashes,'code_sha256':sha256(Path(__file__)),
            'route_selection_counts':tables['routes'].groupby('selection_reason').size().to_dict(),
            'retained_modes':pd.read_sql_query('SELECT mode,count(*) routes FROM routes GROUP BY mode',c).to_dict('records'),
            'environment':{m:importlib.metadata.version(m) for m in ['pandas','numpy','geopandas','shapely','scipy','openpyxl']}}
        for key,value in metadata.items():c.execute('INSERT INTO processing_metadata VALUES (?,?)',(key,json.dumps(value,ensure_ascii=False)))
        c.execute('ANALYZE');c.commit()
    finally:
        c.close()
    with sqlite3.connect(temporary) as saved:
        require(saved.execute('PRAGMA quick_check').fetchone()[0]=='ok','최종 SQLite 재읽기 오류')
        require(saved.execute('SELECT count(*) FROM stop_times').fetchone()[0]==counts['stop_times'],'재읽기 행 수 불일치')
        require(json.loads(saved.execute("SELECT value_json FROM processing_metadata WHERE key='code_sha256'").fetchone()[0])==sha256(Path(__file__)),'저장 코드 해시 불일치')
    saved.close()
    temporary.replace(output)
    save_report('gtfs', metadata)
    print(json.dumps(metadata,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
